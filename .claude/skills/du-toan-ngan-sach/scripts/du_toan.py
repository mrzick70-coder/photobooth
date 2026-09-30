"""Tính dự toán theo ngân sách, tự chọn mức giá, xuất Excel có công thức.

Cách dùng:
    python3 du_toan.py <thu-muc-du-an>

Đầu vào:  <thu-muc-du-an>/hang-muc.json   (do bước thiết kế tạo ra)
          <thu-muc-du-an>/../bang-gia.csv  (bảng giá dùng chung của repo; lần đầu tự
                                            chép từ bang-gia-mau.csv của skill)
Đầu ra:   <thu-muc-du-an>/du-toan.xlsx     + bản tóm tắt ngắn in ra màn hình

Mã thoát: 0 = đạt ngân sách · 1 = vượt ngân sách · 2 = thiếu giá (in danh sách mã cần khảo giá)
"""
import csv
import shutil
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

GIA_MAU = Path(__file__).resolve().parent.parent / "bang-gia-mau.csv"
HAN_GIA_NGAY = 120  # giá cũ hơn số ngày này thì cảnh báo nên khảo lại


def doc_bang_gia(bang_gia):
    if not bang_gia.exists():
        shutil.copy(GIA_MAU, bang_gia)
        print(f"Đã tạo {bang_gia} từ bảng giá mẫu")
    with open(bang_gia, encoding="utf-8") as f:
        return {r["ma"]: r for r in csv.DictReader(f)}


def don_gia(gia, ma, muc):
    return 0 if muc == "KHONG" else int(gia[ma][f"gia_{muc}"])


def chon_muc(items, gia, tran):
    """Tham lam theo độ ưu tiên (1 = lên ảnh nhiều nhất).

    B1: hạng mục bắt buộc ở mức TK (hoặc mức cố định 'muc' nếu có).
    B2: thêm hạng mục tùy chọn ở mức TK nếu còn vừa.
    B3: nâng CB rồi NB (chỉ khi cho_phep_nb) theo thứ tự ưu tiên.
    """
    muc = {}
    for it in items:
        muc[it["ma"]] = it.get("muc") or ("TK" if it.get("bat_buoc", True) else "KHONG")
    tong = lambda: sum(don_gia(gia, it["ma"], muc[it["ma"]]) * it["khoi_luong"] for it in items)
    thu_tu = sorted((it for it in items if not it.get("muc")), key=lambda it: it.get("uu_tien", 3))
    buoc = [("TK", lambda it, m: m == "KHONG"),
            ("CB", lambda it, m: m == "TK"),
            ("NB", lambda it, m: m == "CB" and it.get("cho_phep_nb", False))]
    for cap, dk in buoc:
        for it in thu_tu:
            cu = muc[it["ma"]]
            if dk(it, cu):
                muc[it["ma"]] = cap
                if tong() > tran:
                    muc[it["ma"]] = cu
    return muc, tong()


def xuat_excel(duan, items, gia, muc, ngan_sach, du_phong, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dự toán"
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    blue = Font(name="Arial", size=10, color="0000FF")
    vang = PatternFill("solid", start_color="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0;(#,##0);"-"'

    ws["A1"] = f"DỰ TOÁN: {duan.get('ten', '')}"
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    for i, (lab, val, fmt) in enumerate([("Ngân sách (đ)", ngan_sach, money), ("Dự phòng", du_phong, "0%")], start=3):
        ws.cell(i, 1, lab).font = bold
        c = ws.cell(i, 3, val)
        c.font, c.number_format, c.fill = blue, fmt, vang
    ws["A5"] = "Sửa ô vàng. Cột 'Mức chọn': TK tiết kiệm · CB cân bằng · NB nổi bật · KHONG bỏ. Tổng tự tính lại."
    ws["A5"].font = Font(name="Arial", size=9, italic=True)

    hdr = ["STT", "Nhóm", "Hạng mục", "Đơn vị", "Khối lượng", "Giá TK", "Giá CB", "Giá NB",
           "Mức chọn", "Đơn giá chọn", "Thành tiền", "Ghi chú", "Nguồn giá (ngày)"]
    H = 7
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(H, j, h)
        c.font, c.border, c.fill = bold, box, PatternFill("solid", start_color="EDE6DA")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    dv = DataValidation(type="list", formula1='"TK,CB,NB,KHONG"')
    ws.add_data_validation(dv)

    first = r = H + 1
    for n, it in enumerate(items, start=1):
        g = gia[it["ma"]]
        vals = [n, it.get("nhom", ""), it.get("ten") or g["hang_muc"], g["don_vi"], it["khoi_luong"],
                int(g["gia_TK"]), int(g["gia_CB"]), int(g["gia_NB"]), muc[it["ma"]],
                f'=IF(I{r}="TK",F{r},IF(I{r}="CB",G{r},IF(I{r}="NB",H{r},0)))', f"=E{r}*J{r}",
                it.get("ghi_chu", ""), f'{g["nguon"]} ({g["ngay_cap_nhat"]})']
        for j, v in enumerate(vals, start=1):
            c = ws.cell(r, j, v)
            c.font, c.border = (blue if j in (5, 6, 7, 8, 9) else base), box
            if j in (6, 7, 8, 10, 11):
                c.number_format = money
        ws.cell(r, 9).fill = vang
        dv.add(ws.cell(r, 9))
        r += 1
    last = r - 1

    r += 1
    t = r
    dong = [("Tổng trước dự phòng", f"=SUM(K{first}:K{last})"), ("Dự phòng", f"=K{t}*$C$4"),
            ("TỔNG DỰ TOÁN", f"=K{t}+K{t+1}"), ("Còn dư (+) / vượt (-)", f"=$C$3-K{t+2}"),
            ("Kết luận", f'=IF(K{t+3}>=0,"ĐẠT","VƯỢT")')]
    for i, (lab, f) in enumerate(dong):
        ws.cell(t + i, 3, lab).font = bold
        c = ws.cell(t + i, 11, f)
        c.font, c.number_format = bold, money
    ws.conditional_formatting.add(f"K{t+4}", FormulaRule(formula=[f"K{t+3}>=0"], fill=PatternFill("solid", start_color="C6EFCE")))
    ws.conditional_formatting.add(f"K{t+4}", FormulaRule(formula=[f"K{t+3}<0"], fill=PatternFill("solid", start_color="FFC7CE")))

    g0 = t + 7
    ws.cell(g0, 3, "TỔNG THEO NHÓM").font = bold
    for i, nhom in enumerate(dict.fromkeys(it.get("nhom", "") for it in items), start=1):
        ws.cell(g0 + i, 3, nhom).font = base
        c = ws.cell(g0 + i, 11, f"=SUMIF($B${first}:$B${last},C{g0+i},$K${first}:$K${last})")
        c.font, c.number_format = base, money
        p = ws.cell(g0 + i, 12, f"=IF($K${t+2}=0,0,K{g0+i}/$K${t+2})")
        p.font, p.number_format = base, "0.0%"

    for j, w in enumerate([5, 16, 40, 10, 10, 12, 12, 12, 10, 13, 14, 40, 40], start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = ws.cell(H + 1, 4)
    wb.save(out)


def main():
    thu_muc = Path(sys.argv[1])
    duan = json.loads((thu_muc / "hang-muc.json").read_text(encoding="utf-8"))
    items = duan["hang_muc"]
    ngan_sach, du_phong = duan["ngan_sach"], duan.get("du_phong", 0.10)
    gia = doc_bang_gia(thu_muc.resolve().parent / "bang-gia.csv")

    thieu = [it["ma"] for it in items if it["ma"] not in gia]
    if thieu:
        print("THIẾU GIÁ, cần khảo giá các mã:", ", ".join(thieu))
        sys.exit(2)
    cu = [it["ma"] for it in items
          if (date.today() - date.fromisoformat(gia[it["ma"]]["ngay_cap_nhat"])).days > HAN_GIA_NGAY]

    tran = ngan_sach / (1 + du_phong)
    muc, tong = chon_muc(items, gia, tran)
    tong_dp = tong * (1 + du_phong)
    xuat_excel(duan, items, gia, muc, ngan_sach, du_phong, thu_muc / "du-toan.xlsx")

    tr = lambda x: f"{x / 1e6:,.1f}tr"
    print(f"Tổng {tr(tong_dp)} (gồm dự phòng {du_phong:.0%}) / ngân sách {tr(ngan_sach)} → "
          + ("ĐẠT" if tong_dp <= ngan_sach else f"VƯỢT {tr(tong_dp - ngan_sach)}"))
    print("Mức chọn:", ", ".join(f"{k}={v}" for k, v in muc.items()))
    if cu:
        print(f"Cảnh báo: giá cũ hơn {HAN_GIA_NGAY} ngày:", ", ".join(cu))
    if tong_dp > ngan_sach:
        # Gợi ý cắt: hạng mục đang tính tiền, ít lên ảnh nhất và đắt nhất đứng trước
        ung_vien = sorted((it for it in items if muc[it["ma"]] != "KHONG"),
                          key=lambda it: (-it.get("uu_tien", 3), -don_gia(gia, it["ma"], muc[it["ma"]]) * it["khoi_luong"]))
        print("Ứng viên cắt/thay (ít lên ảnh, đắt trước):")
        for it in ung_vien[:6]:
            print(f"  - {it['ma']}: {tr(don_gia(gia, it['ma'], muc[it['ma']]) * it['khoi_luong'])}"
                  f" (ưu tiên {it.get('uu_tien', 3)}, {'bắt buộc' if it.get('bat_buoc', True) else 'tùy chọn'})")
        sys.exit(1)


if __name__ == "__main__":
    main()
